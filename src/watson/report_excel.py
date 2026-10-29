"""Excel workbook that prints on A4 landscape, one page wide."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from watson import __version__
from watson.classify import BROKEN, INSECURE, REDIRECT, UNVERIFIED
from watson.models import ScanResult
from watson.present import sort_findings, summary_rows
from watson.redact import excel_safe

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
_BODY_FONT = Font(name="Calibri", size=10)
_WRAP = Alignment(wrap_text=True, vertical="top")
_HEADER_ALIGN = Alignment(wrap_text=True, vertical="center")
_SEVERITY_FILLS = {
    "High": PatternFill("solid", fgColor="F8D0D0"),
    "Medium": PatternFill("solid", fgColor="FDE7D0"),
    "Low": PatternFill("solid", fgColor="FFF4CC"),
    "Info": PatternFill("solid", fgColor="E7EEF8"),
}


def _text(value: object) -> str:
    if value is None:
        return ""
    return excel_safe(str(value))


def _write_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[object]], widths: list[int], severity_column: int | None):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append([_text(cell) for cell in row])
    sheet.freeze_panes = "A2"
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=0.4, right=0.4, top=0.55, bottom=0.45, header=0.22, footer=0.22)
    sheet.print_options.horizontalCentered = True
    sheet.oddFooter.left.text = "Watson — passive check, internal use"
    sheet.oddFooter.right.text = "Page &P of &N"
    sheet.sheet_view.showGridLines = False
    last_column = get_column_letter(len(headers))
    last_row = max(sheet.max_row, 1)
    sheet.auto_filter.ref = f"A1:{last_column}{last_row}"
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGN
    sheet.row_dimensions[1].height = 22
    for row in sheet.iter_rows(min_row=2, max_row=last_row, max_col=len(headers)):
        for cell in row:
            cell.font = _BODY_FONT
            cell.alignment = _WRAP
        if severity_column is not None and row:
            label = row[severity_column - 1].value
            if label in _SEVERITY_FILLS:
                row[severity_column - 1].fill = _SEVERITY_FILLS[label]
    sheet.auto_filter.ref = f"A1:{last_column}{last_row}"
    return sheet


def write_excel(result: ScanResult, path: Path) -> None:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    workbook.properties.creator = f"Watson {__version__}"
    workbook.properties.title = "Watson website check"
    workbook.properties.subject = "Passive website check"
    summary = summary_rows(result)
    _write_sheet(
        workbook,
        "Summary",
        ["Item", "Value"],
        [[item, value] for item, value in summary],
        [36, 88],
        None,
    )
    broken = result.links_in(BROKEN)
    unverified = result.links_in(UNVERIFIED)
    redirects = result.links_in(REDIRECT) + result.links_in(INSECURE)
    link_headers = ["URL", "Severity", "Problem", "Status", "Found on", "Link text", "Owner", "Suggested fix"]
    link_widths = [42, 12, 36, 10, 36, 24, 26, 42]

    def link_row(item, problem: str) -> list[object]:
        return [
            item.url,
            item.severity,
            problem,
            item.status if item.status is not None else "",
            item.source_url,
            item.link_text,
            item.owner,
            item.fix,
        ]

    _write_sheet(
        workbook,
        "Broken links",
        link_headers,
        [link_row(item, item.reason) for item in broken],
        link_widths,
        2,
    )
    _write_sheet(
        workbook,
        "Couldn't verify",
        link_headers,
        [link_row(item, item.reason) for item in unverified],
        link_widths,
        2,
    )
    redirect_headers = ["URL", "Severity", "Kind", "Problem", "Status", "Redirects to", "Found on", "Link text", "Owner", "Suggested fix"]
    _write_sheet(
        workbook,
        "Redirects & http",
        redirect_headers,
        [
            [
                item.url,
                item.severity,
                "Soft 404" if item.soft_404 else ("Insecure http" if item.category == INSECURE else "Redirect"),
                item.reason,
                item.status if item.status is not None else "",
                item.location,
                item.source_url,
                item.link_text,
                item.owner,
                item.fix,
            ]
            for item in redirects
        ],
        [36, 12, 16, 36, 10, 32, 32, 22, 24, 36],
        2,
    )
    _write_sheet(
        workbook,
        "Updates & stale content",
        ["Kind", "Severity", "Page", "Date", "Detail", "Owner", "Suggested fix"],
        [
            [item.kind, item.severity, item.page_url, item.observed_date, item.detail, item.owner, item.fix]
            for item in result.updates
        ],
        [22, 12, 40, 16, 55, 26, 42],
        2,
    )
    _write_sheet(
        workbook,
        "Security",
        ["Severity", "Check", "Title", "Evidence", "Owner", "Suggested fix", "Page"],
        [
            [item.severity, item.check, item.title, item.evidence, item.owner, item.fix, item.url]
            for item in sort_findings(result)
        ],
        [12, 24, 42, 55, 28, 46, 36],
        1,
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
