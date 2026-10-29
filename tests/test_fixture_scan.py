"""End-to-end scan of the built-in example site. No live site is contacted."""

import io
import json
import zipfile
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from watson.cli import main, write_reports
from watson.config import Config
from watson.crawler import allowed_by_robots, robots_from_text
from watson.fixture_site import COOKIE_HEADER, DOCUMENT_AUTHOR, PRIVATE_MARKER, serve_fixture
from watson.scan import run_scan

FORBIDDEN = (
    DOCUMENT_AUTHOR,
    COOKIE_HEADER,
    "office@example.com",
    "info@example.com",
    PRIVATE_MARKER,
)


def _report_text(path: Path) -> str:
    data = path.read_bytes()
    if path.suffix == ".xlsx":
        parts = []
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            for name in archive.namelist():
                parts.append(archive.read(name).decode("utf-8", "ignore"))
        return "\n".join(parts)
    return data.decode("utf-8", "ignore")


def test_robots_prefix_blocks_the_draft_page():
    parser = robots_from_text("User-agent: *\nDisallow: /private-drafts\n")
    assert allowed_by_robots(parser, "Watson/0.1.0", "http://127.0.0.1/private-drafts.html") is False
    assert allowed_by_robots(parser, "Watson/0.1.0", "http://127.0.0.1/news.html") is True


def test_fixture_scan_fills_every_report_category(tmp_path):
    with serve_fixture() as site:
        config = Config(
            start_url=site.base + "/",
            max_pages=30,
            max_depth=2,
            rate=100,
            timeout=5,
            output_dir=tmp_path,
            today=date(2026, 10, 29),
        )
        result = run_scan(config, log=lambda _message: None)
        paths = write_reports(result, tmp_path)
        hits = list(site.hits)

    assert result.exit_code == 1
    assert result.links_in("broken")
    assert result.links_in("unverified")
    assert any(item.soft_404 for item in result.links)
    assert result.links_in("insecure")
    assert any(row.kind == "Stale content" and "2024-06-03" in row.detail for row in result.updates)
    assert any("2024" in row.detail and "before" in row.detail for row in result.updates)
    assert any("2019" in row.detail for row in result.updates)
    assert any("2024-08-20" in row.detail for row in result.updates)
    assert not any("2020-02-02" in row.detail for row in result.updates)
    assert result.newest_newsletter == "2024-06-03"
    assert result.newest_news == "2026-10-01"
    checks = {finding.check for finding in result.findings}
    for required in (
        "https-enforcement",
        "content-security-policy",
        "clickjacking",
        "cookie-httponly",
        "version-banner",
        "outdated-library",
        "exposed-email",
        "document-metadata",
        "dangling-domain",
        "insecure-form",
        "security-txt",
        "robots-sitemap",
        "third-party-embed",
    ):
        assert required in checks, required
    assert any(finding.severity == "High" for finding in result.findings)
    assert any(finding.owner == "Platform/hosting provider" for finding in result.findings)
    assert any(finding.owner == "Content editor" for finding in result.findings)
    email = next(finding for finding in result.findings if finding.check == "exposed-email")
    assert "2 unique" in email.evidence

    requested = {(method, path) for method, path, _agent in hits}
    assert ("HEAD", "/only-get") in requested
    assert ("GET", "/only-get") in requested
    assert not any(path.startswith("/private-drafts") for _method, path, _agent in hits)
    assert all("Watson" in agent and "github.com/JeffStone69/Watson" in agent for _m, _p, agent in hits)

    workbook = load_workbook(paths["xlsx"])
    assert workbook.sheetnames == [
        "Summary",
        "Broken links",
        "Couldn't verify",
        "Redirects & http",
        "Updates & stale content",
        "Security",
    ]
    for name in workbook.sheetnames:
        sheet = workbook[name]
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref.startswith("A1:")
        assert sheet.page_setup.fitToWidth == 1
        assert sheet.page_setup.fitToHeight == 0
        assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
        assert sheet.print_title_rows in {"1:1", "$1:$1"}
    assert workbook["Broken links"].max_row > 1
    assert workbook["Security"].max_row > 1

    html = paths["html"].read_text(encoding="utf-8")
    assert "<script" not in html.lower()
    assert "cdnjs" not in html
    assert "fonts.googleapis" not in html
    assert "passive" in html.lower()

    payload = json.loads(paths["json"].read_text(encoding="utf-8"))
    assert payload["passive_only"] is True
    assert payload["exit_code"] == 1
    assert payload["links"]["broken"]
    assert "html" not in payload["pages"][0]

    for path in paths.values():
        text = _report_text(path)
        for secret in FORBIDDEN:
            assert secret not in text, f"{secret} leaked into {path.name}"


def test_cli_demo_writes_three_reports(tmp_path):
    code = main(
        [
            "--demo",
            "--output-dir",
            str(tmp_path),
            "--rate",
            "100",
            "--timeout",
            "5",
            "--max-pages",
            "30",
        ]
    )
    assert code == 1
    assert (tmp_path / "report.xlsx").is_file()
    assert (tmp_path / "report.html").is_file()
    assert (tmp_path / "report.json").is_file()


def test_cli_usage_codes():
    assert main(["--version"]) == 0
    assert main([]) == 2
    assert main(["--demo", "https://www.example.com/"]) == 2
